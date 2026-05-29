"""合同文本脱敏服务。"""

import io
import re
from dataclasses import dataclass

from docx import Document


@dataclass
class SanitizationResult:
    """脱敏结果，包含脱敏文本和映射关系。"""

    sanitized_text: str
    mappings: list[dict]
    errors: list[str]


SENSITIVE_PATTERNS = [
    (
        "company",
        "公司",
        re.compile(
            r"[\u4e00-\u9fa5]{2,50}(?:有限责任公司|科技有限公司|有限公司|集团|银行|公司)"
        ),
    ),
    ("credit_code", "统一社会信用代码", re.compile(r"[0-9A-Z]{18}")),
    ("id_card", "身份证", re.compile(r"\d{17}[\dXx]")),
    ("mobile", "电话", re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")),
    ("telephone", "电话", re.compile(r"(?<!\d)\d{3,4}[-\s]?\d{7,8}(?!\d)")),
    ("email", "邮箱", re.compile(r"[\w.-]+@[\w.-]+\.\w+")),
    ("bank_account", "银行账号", re.compile(r"(?<!\d)(?:\d[ -]?){12,19}(?!\d)")),
    (
        "bank_name",
        "开户行",
        re.compile(r"开户行[:：]?\s*[\u4e00-\u9fa5]{2,40}(?:银行|支行|分行)"),
    ),
    (
        "address",
        "地址",
        re.compile(
            r"(?:地址|住所|注册地址)[:：]?\s*[\u4e00-\u9fa5A-Za-z0-9\-号弄室座层栋路街区县市省]{6,80}"
        ),
    ),
]


def sanitize_contract_text(text: str) -> SanitizationResult:
    """脱敏合同文本，保留金额和日期。"""
    if not text:
        return SanitizationResult(
            sanitized_text="", mappings=[], errors=["合同文本为空，无法脱敏"]
        )

    sanitized = text
    mappings: list[dict] = []
    counters: dict[str, int] = {}
    original_to_placeholder: dict[tuple[str, str], str] = {}

    for sensitive_type, label, pattern in SENSITIVE_PATTERNS:

        def replace(match: re.Match) -> str:
            original = match.group(0)
            key = (sensitive_type, original)
            if key not in original_to_placeholder:
                counters[label] = counters.get(label, 0) + 1
                placeholder = f"{label}{_index_to_letter(counters[label])}"
                original_to_placeholder[key] = placeholder
                mappings.append(
                    {
                        "placeholder": placeholder,
                        "original": original,
                        "type": sensitive_type,
                    }
                )
            return original_to_placeholder[key]

        sanitized = pattern.sub(replace, sanitized)

    sanitized = _sanitize_person_names(
        sanitized, mappings, counters, original_to_placeholder
    )
    return SanitizationResult(sanitized_text=sanitized, mappings=mappings, errors=[])


def restore_text_from_mapping(text: str, mappings: list[dict]) -> str:
    """将脱敏占位符还原为原文。"""
    restored = text or ""
    for item in sorted(
        mappings or [],
        key=lambda value: len(value.get("placeholder", "")),
        reverse=True,
    ):
        placeholder = item.get("placeholder")
        original = item.get("original")
        if placeholder and original:
            restored = restored.replace(placeholder, original)
    return restored


def apply_sanitization_mappings(text: str, mappings: list[dict]) -> str:
    """将原文中的敏感信息替换为脱敏占位符（restore 的反向操作）。"""
    result = text or ""
    for item in sorted(
        mappings or [], key=lambda value: len(value.get("original", "")), reverse=True
    ):
        placeholder = item.get("placeholder")
        original = item.get("original")
        if placeholder and original:
            result = result.replace(original, placeholder)
    return result


def sanitize_docx_bytes(file_content: bytes, mappings: list[dict]) -> bytes:
    """在原 docx 结构上原地脱敏，保留格式，替换敏感文字为占位符。"""
    doc = Document(io.BytesIO(file_content))

    for para in doc.paragraphs:
        for run in para.runs:
            if run.text:
                run.text = apply_sanitization_mappings(run.text, mappings)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    for run in para.runs:
                        if run.text:
                            run.text = apply_sanitization_mappings(run.text, mappings)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _index_to_letter(index: int) -> str:
    """将 1-based 序号转换为 A/B/C 样式。"""
    letters = []
    value = index
    while value > 0:
        value -= 1
        letters.append(chr(ord("A") + value % 26))
        value //= 26
    return "".join(reversed(letters))


def _sanitize_person_names(
    text: str,
    mappings: list[dict],
    counters: dict[str, int],
    original_to_placeholder: dict[tuple[str, str], str],
) -> str:
    """脱敏带有明确标签的联系人姓名。"""
    pattern = re.compile(
        r"(?P<prefix>(?:联系人|法定代表人|授权代表|经办人|签署人|代表)[:：]?\s*)(?P<name>[\u4e00-\u9fa5]{2,4})"
    )

    def replace(match: re.Match) -> str:
        name = match.group("name")
        key = ("person", name)
        if key not in original_to_placeholder:
            counters["人员"] = counters.get("人员", 0) + 1
            placeholder = f"人员{_index_to_letter(counters['人员'])}"
            original_to_placeholder[key] = placeholder
            mappings.append(
                {
                    "placeholder": placeholder,
                    "original": name,
                    "type": "person",
                }
            )
        return f"{match.group('prefix')}{original_to_placeholder[key]}"

    return pattern.sub(replace, text)
