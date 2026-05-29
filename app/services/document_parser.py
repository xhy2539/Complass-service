"""文档解析服务，支持 docx、pdf、txt 格式的文本提取和结构化处理。"""

import io
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import pdfplumber
from docx import Document
from docx.oxml.ns import qn


class DocumentParseError(Exception):
    """文档解析失败异常。"""

    pass


@dataclass
class Paragraph:
    """段落结构，包含位置信息和文本内容。"""

    index: int  # 段落索引
    text: str  # 段落文本
    char_offset_start: int  # 在文档中的字符起始位置
    char_offset_end: int  # 在文档中的字符结束位置
    page_number: Optional[int] = None  # 页码（PDF 有效）
    is_key_clause: bool = False  # 是否是关键条款

    # 段落类型：heading1/heading2/heading3/body/table
    paragraph_type: str = "body"
    # 标题层级：0=正文, 1=一级标题, 2=二级标题, 3=三级标题
    paragraph_level: int = 0


@dataclass
class ParsedDocument:
    """解析后的文档结构。"""

    text: str  # 纯文本内容
    paragraphs: list[Paragraph]  # 段落列表
    sentences: list[dict]  # 句子列表（用于比对）
    char_count: int  # 字符总数
    page_count: Optional[int]  # 页数
    file_name: str  # 原文件名
    file_type: str  # 文件类型
    sanitized_text: str  # 脱敏后的文本（用于 AI 输入）


class DocumentParser:
    """文档解析器，支持 docx、pdf、txt 文件的文本提取和结构化。"""

    SUPPORTED_EXTENSIONS = {".docx", ".pdf", ".txt"}

    # 脱敏规则
    SANITIZE_PATTERNS = [
        (r"\d{11}", "[手机号]"),  # 手机号
        (r"\d{3,4}[-\s]?\d{7,8}", "[电话号码]"),  # 固定电话
        (r"[\w.-]+@[\w.-]+\.\w+", "[邮箱]"),  # 邮箱
        (r"622\d{13}", "[银行卡]"),  # 银行卡号（简单判断）
    ]

    # 标题识别正则表达式（按优先级排序，作为 style 检测的兜底）
    HEADING_PATTERNS = [
        (
            r"^第[一二三四五六七八九十百千零\d]+[条章节款项]",
            "heading2",
        ),  # "第X条"、"第X章"
        (r"^[一二三四五六七八九十百千零\d]+[、.。]", "heading3"),  # "一、"、"1."
        (r"^[《『「]?[一-龥]{2,20}[》』」]?$", "heading1"),  # 纯中文标题，2-20个汉字
        (r"^[一-龥]{1,10}$", "heading3"),  # 短中文文本，可能是小标题
    ]

    # docx 原生标题样式名（优先于正则匹配）
    HEADING_STYLE_PATTERNS = [
        (r"heading\s*1|标题\s*1", "heading1", 1),
        (r"heading\s*2|标题\s*2", "heading2", 2),
        (r"heading\s*3|标题\s*3", "heading3", 3),
    ]

    @classmethod
    def is_supported(cls, filename: str) -> bool:
        """检查文件扩展名是否支持。"""
        ext = Path(filename).suffix.lower()
        return ext in cls.SUPPORTED_EXTENSIONS

    @classmethod
    def get_extension(cls, filename: str) -> str:
        """获取文件扩展名（小写）。"""
        return Path(filename).suffix.lower()

    @classmethod
    def _sanitize_text(cls, text: str) -> str:
        """对文本进行脱敏处理。"""
        result = text
        for pattern, replacement in cls.SANITIZE_PATTERNS:
            result = re.sub(pattern, replacement, result)
        return result

    @classmethod
    def _detect_docx_heading_from_style(cls, para) -> tuple[str, int] | None:
        """从 docx 段落原生样式中检测标题层级，返回 (type, level) 或 None。"""
        try:
            style_name = (para.style.name or "").lower()
        except Exception:
            return None
        for pattern, ptype, level in cls.HEADING_STYLE_PATTERNS:
            if re.search(pattern, style_name):
                return ptype, level
        return None

    @classmethod
    def _detect_paragraph_type(cls, text: str) -> tuple[str, int]:
        """
        检测段落类型和层级（基于文本内容正则匹配）。

        Returns:
            (paragraph_type, paragraph_level)
        """
        text = text.strip()
        if not text:
            return "body", 0

        # 长度异常（过长或过短）可能是正文
        if len(text) > 200 or len(text) < 2:
            return "body", 0

        for pattern, ptype in cls.HEADING_PATTERNS:
            if re.match(pattern, text):
                if ptype == "heading1":
                    return "heading1", 1
                elif ptype == "heading2":
                    return "heading2", 2
                elif ptype == "heading3":
                    return "heading3", 3

        return "body", 0

    @classmethod
    def _add_paragraph_entry(
        cls,
        paragraphs: list[Paragraph],
        sentences: list[dict],
        text: str,
        char_offset: int,
        paragraph_type: str,
        paragraph_level: int,
        filename: str,
        page_number: int | None = None,
    ) -> int:
        """添加一个段落和对应句子，返回 char_offset 增量。"""
        para_start = char_offset
        para_end = char_offset + len(text)

        paragraph = Paragraph(
            index=len(paragraphs),
            text=text,
            char_offset_start=para_start,
            char_offset_end=para_end,
            page_number=page_number,
            is_key_clause=cls._is_key_clause(text),
            paragraph_type=paragraph_type,
            paragraph_level=paragraph_level,
        )
        paragraphs.append(paragraph)

        for sent_match in re.finditer(r"[^。！？;]+[。！？;]*", text):
            sent_text = sent_match.group()
            if sent_text:
                sentences.append(
                    {
                        "text": sent_text,
                        "char_offset_start": para_start + sent_match.start(),
                        "char_offset_end": para_start + sent_match.end(),
                        "paragraph_index": len(paragraphs) - 1,
                        "file_name": filename,
                    }
                )

        return len(text) + 1  # +1 for \n separator

    @classmethod
    def _parse_docx(cls, file_content: bytes, filename: str) -> ParsedDocument:
        """解析 docx 文件，按文档顺序遍历段落和表格。"""
        try:
            doc = Document(io.BytesIO(file_content))
            paragraphs: list[Paragraph] = []
            sentences: list[dict] = []
            char_offset = 0

            body = doc.element.body
            paras_iter = iter(doc.paragraphs)
            tables_iter = iter(doc.tables)

            for child in body:
                if child.tag == qn("w:p"):
                    try:
                        para = next(paras_iter)
                    except StopIteration:
                        continue
                    text = para.text.strip()
                    if not text:
                        continue

                    # 优先从 Word 原生样式检测标题
                    heading = cls._detect_docx_heading_from_style(para)
                    if heading:
                        paragraph_type, paragraph_level = heading
                    else:
                        paragraph_type, paragraph_level = cls._detect_paragraph_type(
                            text
                        )

                    char_offset += cls._add_paragraph_entry(
                        paragraphs,
                        sentences,
                        text,
                        char_offset,
                        paragraph_type,
                        paragraph_level,
                        filename,
                    )

                elif child.tag == qn("w:tbl"):
                    try:
                        table = next(tables_iter)
                    except StopIteration:
                        continue
                    table_rows: list[str] = []
                    for row in table.rows:
                        cell_texts = []
                        for cell in row.cells:
                            ct = cell.text.strip()
                            if ct:
                                cell_texts.append(ct)
                        if cell_texts:
                            table_rows.append(" | ".join(cell_texts))
                    if table_rows:
                        table_text = "【表格】\n" + "\n".join(table_rows)
                        char_offset += cls._add_paragraph_entry(
                            paragraphs,
                            sentences,
                            table_text,
                            char_offset,
                            "body",
                            0,
                            filename,
                        )

            text = "\n".join(p.text for p in paragraphs)

            return ParsedDocument(
                text=text,
                paragraphs=paragraphs,
                sentences=sentences,
                char_count=len(text),
                page_count=None,
                file_name=filename,
                file_type="docx",
                sanitized_text=cls._sanitize_text(text),
            )
        except Exception as e:
            raise DocumentParseError(f"解析 docx 文件失败: {e}")

    @classmethod
    def _parse_pdf(cls, file_content: bytes, filename: str) -> ParsedDocument:
        """解析 pdf 文件。"""
        try:
            paragraphs: list[Paragraph] = []
            sentences: list[dict] = []
            char_offset = 0

            with pdfplumber.open(io.BytesIO(file_content)) as pdf:
                page_count = len(pdf.pages)

                for page_num, page in enumerate(pdf.pages, start=1):
                    text = page.extract_text()
                    if not text:
                        continue

                    lines = text.split("\n")
                    for line in lines:
                        line = line.strip()
                        if not line:
                            continue

                        paragraph_type, paragraph_level = cls._detect_paragraph_type(
                            line
                        )
                        char_offset += cls._add_paragraph_entry(
                            paragraphs,
                            sentences,
                            line,
                            char_offset,
                            paragraph_type,
                            paragraph_level,
                            filename,
                            page_number=page_num,
                        )

            text = "\n".join(p.text for p in paragraphs)

            return ParsedDocument(
                text=text,
                paragraphs=paragraphs,
                sentences=sentences,
                char_count=len(text),
                page_count=page_count,
                file_name=filename,
                file_type="pdf",
                sanitized_text=cls._sanitize_text(text),
            )
        except Exception as e:
            raise DocumentParseError(f"解析 pdf 文件失败: {e}")

    @classmethod
    def _parse_txt(cls, file_content: bytes, filename: str) -> ParsedDocument:
        """解析 txt 文件。"""
        try:
            try:
                text = file_content.decode("utf-8")
            except UnicodeDecodeError:
                text = file_content.decode("gbk")

            paragraphs: list[Paragraph] = []
            sentences: list[dict] = []
            char_offset = 0

            lines = text.split("\n")
            for line in lines:
                line = line.strip()
                if not line:
                    continue

                paragraph_type, paragraph_level = cls._detect_paragraph_type(line)
                char_offset += cls._add_paragraph_entry(
                    paragraphs,
                    sentences,
                    line,
                    char_offset,
                    paragraph_type,
                    paragraph_level,
                    filename,
                )

            joined_text = "\n".join(p.text for p in paragraphs)

            return ParsedDocument(
                text=joined_text,
                paragraphs=paragraphs,
                sentences=sentences,
                char_count=len(joined_text),
                page_count=None,
                file_name=filename,
                file_type="txt",
                sanitized_text=cls._sanitize_text(joined_text),
            )
        except Exception as e:
            raise DocumentParseError(f"解析 txt 文件失败: {e}")

    @classmethod
    def _is_key_clause(cls, text: str) -> bool:
        """判断是否为关键条款。"""
        key_keywords = [
            "违约",
            "赔偿",
            "责任",
            "罚款",
            "解除",
            "终止",
            "付款",
            "金额",
            "交付",
            "保密",
            "知识产权",
        ]
        return any(kw in text for kw in key_keywords)

    @classmethod
    def parse(cls, file_content: bytes, filename: str) -> ParsedDocument:
        """
        解析文档并返回结构化结果。

        Args:
            file_content: 文件二进制内容
            filename: 文件名

        Returns:
            ParsedDocument 对象
        """
        ext = cls.get_extension(filename)

        if ext == ".docx":
            return cls._parse_docx(file_content, filename)
        elif ext == ".pdf":
            return cls._parse_pdf(file_content, filename)
        elif ext == ".txt":
            return cls._parse_txt(file_content, filename)
        else:
            raise DocumentParseError(f"不支持的文件格式: {ext}")
