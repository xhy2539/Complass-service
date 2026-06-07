"""功能性测试：文档解析、脱敏全链路、文档导出。"""

import os

import pytest

# ==================== 文档解析 ====================


def test_document_parser_parse_txt():
    """TXT 文件提取文本。"""
    from app.services.document_parser import DocumentParser

    result = DocumentParser.parse(b"Hello World", "test.txt")
    assert result.text.strip() == "Hello World"
    assert result.file_type == "txt"


def test_document_parser_parse_docx():
    """DOCX 文件提取文本和段落结构。"""
    from app.services.document_parser import DocumentParser

    path = "C:/Users/ASUS/Desktop/01_智能设备采购合同_基准版.docx"
    if not os.path.exists(path):
        pytest.skip("test contract not available")

    with open(path, "rb") as f:
        content = f.read()
    result = DocumentParser.parse(content, "合同.docx")
    assert len(result.text) > 100, "Should extract substantial text"
    assert result.file_type == "docx"
    assert result.char_count > 0
    assert len(result.paragraphs) > 5, "Should identify multiple paragraphs"


def test_document_parser_supported_formats():
    """仅支持 docx/pdf/txt。"""
    from app.services.document_parser import DocumentParser

    for good in ("contract.docx", "file.pdf", "readme.txt"):
        assert DocumentParser.is_supported(good), f"{good} should be supported"
    for bad in ("image.jpg", "video.mp4", "data.xlsx"):
        assert not DocumentParser.is_supported(bad), f"{bad} should be rejected"


# ==================== 脱敏全链路 ====================


def test_sanitize_round_trip_with_real_contract():
    """真实合同：脱敏 → 还原 → 原文一致。"""
    from app.services.document_parser import DocumentParser
    from app.services.sanitization_service import restore_text_from_mapping
    from app.services.sanitization_service import sanitize_contract_text

    path = "C:/Users/ASUS/Desktop/01_智能设备采购合同_基准版.docx"
    if not os.path.exists(path):
        pytest.skip("test contract not available")

    with open(path, "rb") as f:
        result = DocumentParser.parse(f.read(), "合同.docx")

    original = result.text
    san = sanitize_contract_text(original)
    assert len(san.mappings) >= 1, (
        f"Should detect at least 1 sensitive item, got {len(san.mappings)}"
    )

    # 模拟 LLM 返回的占位符文本
    ai_output = san.sanitized_text
    restored = restore_text_from_mapping(ai_output, san.mappings)

    # 还原后不能有占位符
    for m in san.mappings:
        placeholder = m["placeholder"]
        assert placeholder not in restored, f"Placeholder '{placeholder}' not restored"

    # 原文敏感信息应该恢复
    for m in san.mappings:
        original_text = m["original"]
        assert original_text in restored, (
            f"Original '{original_text[:30]}' not found in restored text"
        )


# ==================== DOCX 导出 ====================


def test_docx_export_preserves_paragraph_count():
    """导出后段落数不变。"""

    from docx import Document

    from app.services.document_exporter import DocumentExporter

    path = "C:/Users/ASUS/Desktop/01_智能设备采购合同_基准版.docx"
    if not os.path.exists(path):
        pytest.skip("test contract not available")

    # 构造修改后文本（只改内容，不动段落）
    orig_doc = Document(path)
    orig_paras = [p.text for p in orig_doc.paragraphs]
    new_text = "\n\n".join(orig_paras)
    new_text = new_text.replace("60台", "72台").replace("1,280,000", "1,468,000")

    buf = DocumentExporter.export_text_to_docx_preserve_format(path, new_text)
    exported = Document(buf)

    assert len(exported.paragraphs) == len(orig_doc.paragraphs), (
        f"Paragraph count changed: {len(orig_doc.paragraphs)} -> {len(exported.paragraphs)}"
    )
    assert len(exported.tables) == len(orig_doc.tables), "Tables should be preserved"

    # 验证修改生效
    body_text = "\n".join(p.text for p in exported.paragraphs)
    assert "72台" in body_text, "Edited text should appear in exported doc"


def test_docx_export_preserves_style_and_bold():
    """导出后段落样式和加粗保留。"""

    from docx import Document

    from app.services.document_exporter import DocumentExporter

    path = "C:/Users/ASUS/Desktop/01_智能设备采购合同_基准版.docx"
    if not os.path.exists(path):
        pytest.skip("test contract not available")

    orig_doc = Document(path)
    orig_paras = [p.text for p in orig_doc.paragraphs]
    new_text = "\n\n".join(orig_paras)

    buf = DocumentExporter.export_text_to_docx_preserve_format(path, new_text)
    exported = Document(buf)

    mismatches = 0
    for i in range(len(orig_doc.paragraphs)):
        op = orig_doc.paragraphs[i]
        ep = exported.paragraphs[i]
        if op.style.name != ep.style.name:
            mismatches += 1
        if op.alignment != ep.alignment:
            mismatches += 1
        for j in range(min(len(op.runs), len(ep.runs))):
            if op.runs[j].bold != ep.runs[j].bold:
                mismatches += 1

    assert mismatches == 0, f"Found {mismatches} format mismatches in exported document"
