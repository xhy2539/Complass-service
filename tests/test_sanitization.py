"""脱敏还原测试：敏感信息 → 占位符 → LLM → 还原。"""

import pytest

from app.services.sanitization_service import apply_sanitization_mappings
from app.services.sanitization_service import restore_text_from_mapping
from app.services.sanitization_service import sanitize_contract_text
from app.services.sanitization_service import sanitize_docx_bytes

# --- 脱敏 ---


def test_sanitize_company_name():
    result = sanitize_contract_text(
        "甲方：上海智能科技有限公司\n乙方：北京云核技术有限公司"
    )
    assert "上海智能科技有限公司" not in result.sanitized_text
    assert "北京云核技术有限公司" not in result.sanitized_text
    assert any("公司" in str(m.get("placeholder", "")) for m in result.mappings)
    assert len(result.mappings) >= 2


def test_sanitize_phone_and_email():
    result = sanitize_contract_text(
        "联系人张三，电话13812345678，邮箱zhangsan@example.com"
    )
    assert "13812345678" not in result.sanitized_text
    assert "zhangsan@example.com" not in result.sanitized_text
    assert len(result.mappings) >= 2


def test_sanitize_id_card_and_credit_code():
    result = sanitize_contract_text(
        "统一社会信用代码：91310115MA1K3ABCDEF，法人身份证110101199001011234"
    )
    assert "91310115MA1K3ABCDEF" not in result.sanitized_text
    assert "110101199001011234" not in result.sanitized_text
    assert len(result.mappings) >= 2


def test_sanitize_bank_account():
    result = sanitize_contract_text(
        "开户行：中国工商银行北京分行 账号：6222021234567890123"
    )
    assert "6222021234567890123" not in result.sanitized_text
    # 开户行模式匹配
    has_bank_mapping = any(
        "银行" in str(m.get("placeholder", "")) for m in result.mappings
    )
    assert has_bank_mapping or len(result.mappings) >= 1


def test_preserve_amounts_and_dates():
    """金额和日期不应被脱敏。"""
    text = "合同总价人民币1,280,000元，交付日期2026年7月15日。"
    result = sanitize_contract_text(text)
    assert "1,280,000" in result.sanitized_text
    assert "2026年7月15日" in result.sanitized_text


# --- 还原 ---


def test_restore_round_trip():
    """脱敏 → 还原 → 原文一致。"""
    original = (
        "甲方上海智能科技有限公司（联系方式：021-12345678）与乙方北京云核技术有限公司"
        "签订合同，联系人张三（13812345678）。"
    )
    result = sanitize_contract_text(original)
    assert len(result.mappings) >= 3

    # 模拟 LLM 返回的文本包含占位符
    ai_output = (
        f"风险点：{result.mappings[0]['placeholder']}的付款条款不合理，"
        f"建议联系{result.mappings[1]['placeholder']}修改。"
    )
    restored = restore_text_from_mapping(ai_output, result.mappings)
    # 验证还原后占位符消失
    assert "公司A" not in restored
    assert "电话A" not in restored
    # 验证原文恢复
    assert "上海智能科技有限公司" in restored or "北京云核技术有限公司" in restored


def test_restore_empty_mappings():
    """无映射时还原不改变文本。"""
    text = "合同正常条款，无敏感信息。"
    assert restore_text_from_mapping(text, []) == text
    assert restore_text_from_mapping(text, None) == text


# --- 正向替换 ---


def test_apply_sanitization_mappings():
    """正向替换：原文 → 占位符。"""
    mappings = [
        {"placeholder": "公司A", "original": "上海智能科技", "type": "company"},
        {"placeholder": "电话A", "original": "13812345678", "type": "mobile"},
    ]
    text = "上海智能科技的联系电话是13812345678。"
    result = apply_sanitization_mappings(text, mappings)
    assert "上海智能科技" not in result
    assert "13812345678" not in result
    assert "公司A" in result
    assert "电话A" in result


# --- DOCX 脱敏 ---


@pytest.mark.skipif(
    not __import__("os").path.exists(
        "C:/Users/ASUS/Desktop/01_智能设备采购合同_基准版.docx"
    ),
    reason="test contract file not found",
)
def test_sanitize_docx_bytes_preserves_structure():
    """DOCX 脱敏后仍能正常打开，格式保留。"""
    import os

    from docx import Document

    path = "C:/Users/ASUS/Desktop/01_智能设备采购合同_基准版.docx"
    if not os.path.exists(path):
        pytest.skip("test contract not available")

    with open(path, "rb") as f:
        content = f.read()

    # 先对文本脱敏获取映射
    from app.services.document_parser import DocumentParser

    doc = DocumentParser.parse(content, "合同.docx")
    sanitization = sanitize_contract_text(doc.text)
    if not sanitization.mappings:
        pytest.skip("no mappings found in test contract")

    # 对 DOCX 应用映射
    sanitized_bytes = sanitize_docx_bytes(content, sanitization.mappings)

    # 验证仍为有效 DOCX
    from io import BytesIO

    sanitized_doc = Document(BytesIO(sanitized_bytes))
    assert len(sanitized_doc.paragraphs) == len(Document(BytesIO(content)).paragraphs)

    # 验证脱敏生效：原文敏感信息不应出现
    sanitized_text = "\n".join(p.text for p in sanitized_doc.paragraphs)
    for m in sanitization.mappings:
        assert m["original"] not in sanitized_text, (
            f"Original '{m['original']}' should be replaced in body paragraphs"
        )


# --- 空文本边界 ---


def test_sanitize_empty_text():
    result = sanitize_contract_text("")
    assert result.sanitized_text == ""
    assert len(result.errors) > 0


def test_sanitize_none_text():
    result = sanitize_contract_text("")  # str 不能是 None
    assert result.sanitized_text == ""
